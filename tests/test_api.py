from __future__ import annotations

import os
import tempfile
import sys
import unittest
import uuid
from unittest.mock import patch
from io import BytesIO
from pathlib import Path

DB_FILE = Path(tempfile.gettempdir()) / f"careerinsight-test-{uuid.uuid4().hex}.sqlite3"
os.environ["DATABASE_URL"] = f"sqlite:///{DB_FILE}"
os.environ["COOKIE_SECURE"] = "false"
os.environ["COOKIE_SAMESITE"] = "lax"
os.environ["RESEARCH_ADMIN_EMAILS"] = "research@example.com"
os.environ.pop("SMTP_HOST", None)
os.environ.pop("SMTP_FROM", None)
os.environ.pop("PUBLIC_APP_URL", None)
os.environ.pop("MANUS_API_URL", None)
os.environ.pop("MANUS_API_KEY", None)
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from fastapi.testclient import TestClient  # noqa: E402
from fastapi import HTTPException  # noqa: E402
from starlette.requests import Request  # noqa: E402
from docx import Document  # noqa: E402
from backend.db import Base, SessionLocal, engine, normalize_database_url  # noqa: E402
from backend.main import app, schema_autocreate_enabled  # noqa: E402
from backend.models import Resume  # noqa: E402
from backend.security import enforce_rate_limit  # noqa: E402


class CareerInsightApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        Base.metadata.create_all(bind=engine)

    def test_schema_autocreation_is_opt_in(self):
        with patch.dict(os.environ):
            os.environ.pop("CAREERINSIGHT_AUTO_CREATE", None)
            self.assertFalse(schema_autocreate_enabled())
            os.environ["CAREERINSIGHT_AUTO_CREATE"] = "true"
            self.assertTrue(schema_autocreate_enabled())

    def test_sensitive_action_rate_limit_returns_retry_after(self):
        request = Request({"type": "http", "client": ("127.0.0.1", 8000), "headers": []})
        subject = uuid.uuid4().hex
        enforce_rate_limit(request, "unit-test", subject, 1, 60)
        with self.assertRaises(HTTPException) as caught:
            enforce_rate_limit(request, "unit-test", subject, 1, 60)
        self.assertEqual(caught.exception.status_code, 429)
        self.assertIn("Retry-After", caught.exception.headers or {})

    def test_managed_database_ssl_url_is_normalized_for_pymysql(self):
        url, connect_args = normalize_database_url("mysql://user:pass@localhost/app?ssl=%7B%22rejectUnauthorized%22%3Atrue%7D")
        self.assertTrue(url.startswith("mysql+pymysql://"))
        self.assertNotIn("ssl=", url)
        self.assertIs(connect_args["ssl"]["verify_mode"], True)
        _, unverified_args = normalize_database_url("mysql://user:pass@localhost/app?ssl=%7B%22rejectUnauthorized%22%3Afalse%7D")
        self.assertIs(unverified_args["ssl"]["verify_mode"], False)

    def setUp(self):
        self.client = TestClient(app)
        self.client.__enter__()
        self.client.get("/api/auth/me")

    def tearDown(self):
        self.client.__exit__(None, None, None)

    def mutate(self, method: str, path: str, **kwargs):
        headers = dict(kwargs.pop("headers", {}))
        headers["X-CSRF-Token"] = self.client.cookies.get("csrf_token", "")
        return self.client.request(method, path, headers=headers, **kwargs)

    def register(self, email: str | None = None):
        address = email or f"learner-{uuid.uuid4().hex[:8]}@example.com"
        response = self.mutate("POST", "/api/auth/register", json={
            "name": "Test Learner", "email": address, "password": "Long-Strong-Password-428!",
            "education": "Undergraduate", "degree": "B.Sc. Computer Science",
            "branch": "Computer Science", "graduation_year": 2026,
        })
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()["user"]

    def test_health_and_csrf_protected_account_profile_and_scoring_flow(self):
        self.assertEqual(self.client.get("/api/healthz").status_code, 200)
        user = self.register()
        self.assertEqual(user["profile_complete"], False)
        self.assertIsNone(self.client.get("/api/dashboard").json()["roadmap"])
        login = self.mutate("POST", "/api/auth/login", json={"email": user["email"], "password": "Long-Strong-Password-428!", "remember": True})
        self.assertEqual(login.status_code, 200, login.text)
        self.assertEqual(login.json()["user"]["id"], user["id"])
        profile = {
            "education": "Undergraduate", "degree": "B.Sc. Computer Science",
            "branch": "Computer Science", "graduation_year": 2026,
            "experience_level": "intermediate", "target_role": "Data Analyst",
            "preferred_industry": "Analytics",
            "skills": [
                {"name": "Python", "level": "advanced"},
                {"name": "SQL", "level": "intermediate"},
                {"name": "Statistics", "level": "intermediate"},
                {"name": "Data Visualization", "level": "beginner"},
            ],
            "soft_skills": ["Communication", "Problem Solving"],
            "projects": [{"name": "Customer analysis", "description": "Analyzed a small practice dataset.", "technologies": ["Python", "SQL"], "domain": "Analytics"}],
            "certifications": [{"name": "Data foundations", "platform": "Sample", "year": 2025}],
        }
        saved = self.mutate("PUT", "/api/profile", json=profile)
        self.assertEqual(saved.status_code, 200, saved.text)
        self.assertEqual(saved.json()["skills"][0]["name"], "Python")
        self.assertTrue(self.client.get("/api/auth/me").json()["user"]["profile_complete"])
        selected = self.mutate("POST", "/api/interests", json={"interests": ["Working with data", "Research"]})
        self.assertEqual(selected.status_code, 200, selected.text)
        self.assertEqual({item["name"] for item in selected.json()["selected"]}, {"Working with data", "Research"})
        invalid_interest = self.mutate("POST", "/api/interests", json={"interests": ["invented high-paying role"]})
        self.assertEqual(invalid_interest.status_code, 422, invalid_interest.text)
        self.assertEqual({item["name"] for item in self.client.get("/api/interests").json()["selected"]}, {"Working with data", "Research"})
        analysis = self.mutate("POST", "/api/career/analyze", json={})
        self.assertEqual(analysis.status_code, 200, analysis.text)
        recommendations = analysis.json()["recommendations"]
        self.assertEqual(len(recommendations), 10)
        self.assertTrue(all(0 <= row["score"] <= 100 for row in recommendations))
        self.assertTrue(all(row["skill_progress"] for row in recommendations))
        self.assertTrue(all(len(row["factors"]) == 6 for row in recommendations))
        self.assertEqual(recommendations[0]["career"], "Data Analyst")
        detail = self.client.get(f"/api/career/{recommendations[0]['career_id']}")
        self.assertEqual(detail.status_code, 200, detail.text)
        gaps = self.client.get("/api/skills/gaps")
        self.assertEqual(gaps.status_code, 200, gaps.text)
        self.assertTrue(gaps.json()["gaps"])
        dashboard = self.client.get("/api/dashboard")
        self.assertEqual(dashboard.status_code, 200, dashboard.text)
        coach = self.mutate("POST", "/api/ai/chat", json={"message": "Compare my top two careers"})
        self.assertEqual(coach.status_code, 200, coach.text)
        self.assertFalse(coach.json()["ai_generated"])
        self.assertIn("%", coach.json()["reply"])
        self.assertTrue(coach.json()["resources"])
        history = self.client.get("/api/ai/history")
        self.assertEqual([item["role"] for item in history.json()["messages"]], ["user", "assistant"])
        self.assertTrue(history.json()["messages"][1]["resources"])
        self.assertEqual(self.client.request("DELETE", "/api/ai/history", json={}).status_code, 403)
        cleared = self.mutate("DELETE", "/api/ai/history", json={})
        self.assertEqual(cleared.status_code, 200, cleared.text)
        self.assertEqual(cleared.json()["deleted"], 2)
        self.assertEqual(self.client.get("/api/ai/history").json()["messages"], [])
        roadmap = dashboard.json()["roadmap"]
        self.assertEqual({item["phase"] for item in roadmap["items"]}, {"Foundation", "Core Skills", "Projects", "Portfolio", "Career Preparation"})
        self.assertTrue(any(item["resource"] for item in roadmap["items"]))
        item_id = roadmap["items"][0]["id"]
        progress = self.mutate("PUT", "/api/roadmap/progress", json={"item_id": item_id, "status": "complete"})
        self.assertEqual(progress.status_code, 200, progress.text)
        self.assertEqual(progress.json()["status"], "complete")
        self.assertTrue(self.client.get("/api/dashboard").json()["roadmap"]["items"][0]["status"] == "complete")
        self.assertEqual(self.client.get("/api/research/evaluations").status_code, 403)
        forgot = self.mutate("POST", "/api/auth/forgot-password", json={"email": user["email"]})
        self.assertEqual(forgot.status_code, 200, forgot.text)
        self.assertNotIn("token", forgot.text.lower())
        invalid_file = self.mutate("POST", "/api/resume/upload", files={"file": ("resume.txt", b"not a resume", "text/plain")})
        self.assertEqual(invalid_file.status_code, 415, invalid_file.text)
        no_csrf_upload = self.client.post("/api/resume/upload", files={"file": ("resume.docx", b"not a resume", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        self.assertEqual(no_csrf_upload.status_code, 403, no_csrf_upload.text)
        document = Document()
        document.add_paragraph("Bachelor of Science in Computer Science")
        document.add_paragraph("Skills: Python, FastAPI, SQL")
        document.add_paragraph("Professional Experience")
        document.add_paragraph("Student analyst at Example Lab, 2024 to 2025")
        document.add_paragraph("Project: Built a research API for career insights")
        document.add_paragraph("Certification: Cloud foundations")
        stream = BytesIO()
        document.save(stream)
        upload = self.mutate("POST", "/api/resume/upload", files={"file": ("review.docx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        self.assertEqual(upload.status_code, 200, upload.text)
        draft = upload.json()
        edited = draft["extracted"] | {"skills": ["Python", "FastAPI"], "education": ["Reviewed degree detail"], "projects": ["Reviewed project detail"], "certifications": ["Reviewed cloud certification"], "experience_excerpts": ["Reviewed student analyst role at Example Lab, 2024 to 2025"]}
        rejected_edit = self.mutate("POST", "/api/resume/accept", json={"draft_id": draft["draft_id"], "extracted": edited | {"skills": "not a list"}})
        self.assertEqual(rejected_edit.status_code, 422)
        accepted = self.mutate("POST", "/api/resume/accept", json={"draft_id": draft["draft_id"], "extracted": edited})
        self.assertEqual(accepted.status_code, 200, accepted.text)
        updated = self.client.get("/api/profile").json()
        self.assertIn("FastAPI", {skill["name"] for skill in updated["skills"]})
        self.assertIn("Reviewed project detail", {project["description"] for project in updated["projects"]})
        self.assertIn("Reviewed cloud certification", {cert["name"] for cert in updated["certifications"]})
        self.assertEqual(updated["experience_level"], "intermediate")
        with SessionLocal() as db:
            saved_resume = db.query(Resume).filter(Resume.user_id == user["id"]).one()
            self.assertEqual(saved_resume.extracted["experience_excerpts"], edited["experience_excerpts"])
            self.assertNotIn("experience_years_mentions", saved_resume.extracted)
        logout = self.mutate("POST", "/api/auth/logout", json={})
        self.assertEqual(logout.status_code, 200, logout.text)
        self.assertIsNone(self.client.get("/api/auth/me").json()["user"])

    def test_researcher_weights_and_evaluation_use_submitted_labels(self):
        self.register("research@example.com")
        changed = self.mutate("PUT", "/api/settings/scoring", json={"weights": {"skills": 40, "interests": 25, "projects": 15, "education": 10, "goal": 5, "experience": 5}})
        self.assertEqual(changed.status_code, 200, changed.text)
        self.assertEqual(sum(changed.json()["weights"].values()), 100)
        rejected = self.mutate("PUT", "/api/settings/scoring", json={"weights": {"skills": 40, "interests": 25, "projects": 15, "education": 10, "goal": 5, "experience": 2}})
        self.assertEqual(rejected.status_code, 422)
        methods = ("baseline", "ml", "ml_llm", "ml_rag_llm")
        cases = [
            {"label": "data-analyst", "predictions": {name: ["data-analyst", "data-scientist"] for name in methods}, "measurements": {name: {"satisfaction": 4, "explanation_usefulness": 3.5, "response_time_ms": 100 + index} for index, name in enumerate(methods)}},
            {"label": "data-scientist", "predictions": {name: ["data-scientist", "data-analyst"] for name in methods}},
        ]
        result = self.mutate("POST", "/api/research/evaluations", json={"dataset_name": "temporary integration test", "cases": cases})
        self.assertEqual(result.status_code, 200, result.text)
        runs = result.json()["runs"]
        self.assertEqual(len(runs), 4)
        self.assertTrue(all(run["sample_count"] == 2 and run["metrics"]["top_3_accuracy"] == 1 for run in runs))
        self.assertTrue(all(run["metrics"]["satisfaction_mean_1_to_5"] == 4 and run["metrics"]["satisfaction_sample_count"] == 1 for run in runs))
        self.assertTrue(all(run["metrics"]["explanation_usefulness_mean_1_to_5"] == 3.5 and run["metrics"]["explanation_usefulness_sample_count"] == 1 for run in runs))
        self.assertTrue(all(run["metrics"]["response_time_sample_count"] == 1 and run["metrics"]["response_time_ms_mean"] >= 100 for run in runs))
        invalid_label = {"dataset_name": "invalid career label", "cases": [{"label": "unknown role", "predictions": {name: ["data-analyst"] for name in methods}}]}
        self.assertEqual(self.mutate("POST", "/api/research/evaluations", json=invalid_label).status_code, 422)
        self.assertEqual(len(self.client.get("/api/research/evaluations").json()["runs"]), 4)

    def test_unknown_experience_and_role_specific_education_and_goal_evidence(self):
        self.register()
        profile = {
            "education": "Undergraduate", "degree": "B.B.A. Business Administration",
            "branch": "Business Administration", "graduation_year": 2026,
            "experience_level": None, "target_role": "software engineer",
            "preferred_industry": "", "skills": [], "soft_skills": [],
            "projects": [], "certifications": [],
        }
        saved = self.mutate("PUT", "/api/profile", json=profile)
        self.assertEqual(saved.status_code, 200, saved.text)
        self.assertIsNone(self.client.get("/api/profile").json()["experience_level"])
        result = self.mutate("POST", "/api/career/analyze", json={})
        self.assertEqual(result.status_code, 200, result.text)
        careers = {row["career"]: row for row in result.json()["recommendations"]}
        data_analyst = {factor["key"]: factor for factor in careers["Data Analyst"]["factors"]}
        developer = {factor["key"]: factor for factor in careers["Software Developer"]["factors"]}
        self.assertEqual(data_analyst["education"]["score"], 100)
        self.assertEqual(developer["education"]["score"], 50)
        self.assertEqual(data_analyst["goal"]["score"], 50)
        self.assertEqual(developer["goal"]["score"], 100)
        self.assertEqual(data_analyst["experience"]["score"], 50)
        self.assertEqual(developer["experience"]["score"], 50)

    def test_catalogue_only_interest_and_specific_project_alignment(self):
        self.register()
        profile = {
            "education": "", "degree": "", "branch": "", "graduation_year": None,
            "experience_level": None, "target_role": "", "preferred_industry": "",
            "skills": [], "soft_skills": [],
            "projects": [{"name": "General exercise", "description": "A research exercise involving data and technology.", "technologies": [], "domain": "General"}],
            "certifications": [],
        }
        self.assertEqual(self.mutate("PUT", "/api/profile", json=profile).status_code, 200)
        self.assertEqual(self.mutate("POST", "/api/interests", json={"interests": ["Product design"]}).status_code, 200)
        analysis = self.mutate("POST", "/api/career/analyze", json={})
        self.assertEqual(analysis.status_code, 200, analysis.text)
        analyst = next(row for row in analysis.json()["recommendations"] if row["career"] == "Data Analyst")
        factors = {factor["key"]: factor for factor in analyst["factors"]}
        self.assertEqual(factors["interests"]["score"], 0)
        self.assertEqual(factors["projects"]["score"], 50)

    def test_resume_review_preserves_explicit_empty_fields_and_experience_excerpts(self):
        user = self.register()
        document = Document()
        document.add_paragraph("Professional Experience")
        document.add_paragraph("Junior analyst at Example Labs, 2024 to 2025")
        document.add_paragraph("Technical Skills")
        document.add_paragraph("Python SQL")
        stream = BytesIO()
        document.save(stream)
        upload = self.mutate("POST", "/api/resume/upload", files={"file": ("experience.docx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        self.assertEqual(upload.status_code, 200, upload.text)
        draft = upload.json()
        self.assertEqual(draft["extracted"]["experience_excerpts"], ["Junior analyst at Example Labs, 2024 to 2025"])
        self.assertNotIn("experience_years_mentions", draft["extracted"])
        accepted_payload = dict(draft["extracted"])
        accepted_payload.pop("skills", None)
        accepted_payload["experience_excerpts"] = ["Reviewed: Junior analyst at Example Labs (2024–2025)."]
        accepted = self.mutate("POST", "/api/resume/accept", json={"draft_id": draft["draft_id"], "extracted": accepted_payload})
        self.assertEqual(accepted.status_code, 200, accepted.text)
        profile = self.client.get("/api/profile").json()
        self.assertNotIn("Python", {skill["name"] for skill in profile["skills"]})
        self.assertIsNone(profile["experience_level"])
        with SessionLocal() as db:
            saved = db.query(Resume).filter(Resume.user_id == user["id"]).one()
            self.assertEqual(saved.extracted["skills"], [])
            self.assertEqual(saved.extracted["experience_excerpts"], accepted_payload["experience_excerpts"])

    def test_auth_me_recovers_missing_csrf_cookie_without_authorizing_a_stale_session(self):
        self.register()
        self.client.cookies.delete("csrf_token")
        response = self.client.get("/api/auth/me")
        self.assertEqual(response.status_code, 200)
        self.assertIsNone(response.json()["user"])
        self.assertTrue(self.client.cookies.get("csrf_token"))
        self.assertIsNone(self.client.cookies.get("career_session"))


if __name__ == "__main__":
    unittest.main(verbosity=2)
